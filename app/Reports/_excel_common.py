"""
Reports/_excel_common.py

Shared OpenPyXL styling helpers (green header, auto-width, freeze panes)
reused by excel_report.py and admin_excel_report.py.
"""

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill(start_color="2E7D32", end_color="2E7D32", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(color="1B5E20", bold=True, size=16)
LABEL_FONT = Font(color="1B5E20", bold=True, size=11)
THIN_BORDER = Border(
    left=Side(style="thin", color="C8E6C9"),
    right=Side(style="thin", color="C8E6C9"),
    top=Side(style="thin", color="C8E6C9"),
    bottom=Side(style="thin", color="C8E6C9"),
)


def write_title(sheet, text, cell="A1"):
    sheet[cell] = text
    sheet[cell].font = TITLE_FONT


def write_header_row(sheet, row_index, headers):
    """Write a styled, bold, green header row and freeze it."""
    for col_index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=row_index, column=col_index, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER
    sheet.freeze_panes = sheet.cell(row=row_index + 1, column=1).coordinate


def write_data_rows(sheet, start_row, rows):
    """Write plain bordered data rows starting at start_row."""
    for r_offset, row in enumerate(rows):
        for c_offset, value in enumerate(row, start=1):
            cell = sheet.cell(row=start_row + r_offset, column=c_offset, value=value)
            cell.border = THIN_BORDER


def auto_size_columns(sheet, min_width=10, max_width=40):
    """Auto-size every column in a sheet based on its longest cell value."""
    widths = {}
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            col_letter = get_column_letter(cell.column)
            widths[col_letter] = max(widths.get(col_letter, 0), len(str(cell.value)) + 2)
    for col_letter, width in widths.items():
        sheet.column_dimensions[col_letter].width = max(min_width, min(width, max_width))
