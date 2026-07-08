"""
Reports/excel_report.py

Builds a two-sheet Excel workbook (Predictions + Summary) for a single user.
"""

from io import BytesIO

from openpyxl import Workbook

from Reports.schema import UserReportData
from Reports._excel_common import (
    write_title, write_header_row, write_data_rows, auto_size_columns, LABEL_FONT,
)
from Reports.utils import format_food_name

PREDICTION_HEADERS = [
    "Date", "Food", "Weight (g)", "Calories", "Protein",
    "Fat", "Carbohydrates", "Fiber", "Sugar", "Accuracy",
]


def _build_predictions_sheet(sheet, data: UserReportData):
    write_title(sheet, "NutriLens - Meal Scans")
    write_header_row(sheet, row_index=3, headers=PREDICTION_HEADERS)

    rows = [
        [
            p.created_at.strftime("%d-%m-%Y %H:%M"),
            format_food_name(p.food_name),
            p.weight_grams,
            p.calories,
            p.protein,
            p.fat,
            p.carbohydrates,
            p.fiber,
            p.sugar,
            p.confidence,
        ]
        for p in data.predictions
    ]
    write_data_rows(sheet, start_row=4, rows=rows)
    auto_size_columns(sheet)


def _write_summary_block(sheet, start_row, title, summary):
    sheet.cell(row=start_row, column=1, value=title).font = LABEL_FONT
    fields = [
        ("Foods Consumed", summary.foods_consumed),
        ("Calories", summary.total_calories),
        ("Protein", summary.total_protein),
        ("Fat", summary.total_fat),
        ("Carbohydrates", summary.total_carbohydrates),
        ("Fiber", summary.total_fiber),
        ("Sugar", summary.total_sugar),
    ]
    for offset, (label, value) in enumerate(fields, start=1):
        sheet.cell(row=start_row + offset, column=1, value=label)
        sheet.cell(row=start_row + offset, column=2, value=value)
    return start_row + len(fields) + 2


def _build_summary_sheet(sheet, data: UserReportData):
    write_title(sheet, f"Nutrition Summary - {data.user.name}")
    row = 3
    row = _write_summary_block(sheet, row, "Today's Summary", data.today)
    row = _write_summary_block(sheet, row, "Weekly Summary", data.weekly)
    row = _write_summary_block(sheet, row, "Monthly Summary", data.monthly)

    sheet.cell(row=row, column=1, value="Top Foods").font = LABEL_FONT
    row += 1
    write_header_row(sheet, row_index=row, headers=["Food", "Times Logged"])
    write_data_rows(sheet, row + 1, [[format_food_name(f.food_name), f.count] for f in data.top_foods])
    auto_size_columns(sheet)


def build_user_excel(data: UserReportData) -> BytesIO:
    """Render a user's nutrition workbook and return it as a BytesIO buffer."""
    workbook = Workbook()

    predictions_sheet = workbook.active
    predictions_sheet.title = "Predictions"
    _build_predictions_sheet(predictions_sheet, data)

    summary_sheet = workbook.create_sheet("Summary")
    _build_summary_sheet(summary_sheet, data)

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer
